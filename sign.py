# property_mail_agent.py
# Reads only new UNSEEN emails, classifies with LLM, replies, and forwards property mails.
# For property mails, attaches a .docx and a .pdf with an interactive signature field.

import os, re, ssl, time, json, email, imaplib, smtplib, traceback, tempfile
from io import BytesIO
from datetime import datetime, timezone
from typing import Optional, Tuple, Dict, Any, List
from email.header import decode_header, make_header
from email.message import EmailMessage
from bs4 import BeautifulSoup
from docx import Document
from dotenv import load_dotenv
from groq import Groq

# PDF libs
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from pdfrw import PdfReader, PdfWriter, IndirectPdfDict

load_dotenv()

# ============== Config from .env ==============
IMAP_HOST        = os.getenv("IMAP_HOST", "")
IMAP_PORT        = int(os.getenv("IMAP_PORT", "993"))
IMAP_USER        = os.getenv("IMAP_USER", "")
IMAP_PASSWORD    = os.getenv("IMAP_PASSWORD", "")
POLL_MAILBOX     = os.getenv("POLL_MAILBOX", "INBOX")

SMTP_HOST        = os.getenv("SMTP_HOST", "")
SMTP_PORT        = int(os.getenv("SMTP_PORT", "465"))
SMTP_USER        = os.getenv("SMTP_USER", "")
SMTP_PASSWORD    = os.getenv("SMTP_PASSWORD", "")
SMTP_FROM_NAME   = os.getenv("SMTP_FROM_NAME", "Mail Agent")

FORWARD_TO_EMAIL = os.getenv("FORWARD_TO_EMAIL", "")
MAX_EMAILS_PER_RUN = int(os.getenv("MAX_EMAILS_PER_RUN", "20"))
DRY_RUN          = os.getenv("DRY_RUN", "false").lower() == "true"
ATTACH_ORIGINAL_EML = os.getenv("ATTACH_ORIGINAL_EML", "false").lower() == "true"

GROQ_API_KEY     = os.getenv("GROQ_API_KEY", "")
GROQ_MODEL       = os.getenv("GROQ_MODEL", "llama-3.1-8b-instant")

for k, v in {
    "IMAP_HOST": IMAP_HOST,
    "IMAP_USER": IMAP_USER,
    "IMAP_PASSWORD": IMAP_PASSWORD,
    "SMTP_HOST": SMTP_HOST,
    "SMTP_USER": SMTP_USER,
    "SMTP_PASSWORD": SMTP_PASSWORD,
    "GROQ_API_KEY": GROQ_API_KEY,
}.items():
    if not v:
        raise SystemExit(f"Missing {k} in .env")

groq = Groq(api_key=GROQ_API_KEY)

# ============== Helpers ==============
def decode_mime(s: Optional[str]) -> str:
    if not s: return ""
    try: return str(make_header(decode_header(s)))
    except Exception: return s

def html_to_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style"]): t.extract()
    return re.sub(r"\n{3,}", "\n\n", soup.get_text("\n")).strip()

def extract_text(msg: email.message.Message) -> Tuple[str, str]:
    plain, html = "", ""
    if msg.is_multipart():
        for part in msg.walk():
            ctype = part.get_content_type()
            disp = str(part.get("Content-Disposition") or "")
            if "attachment" in disp: continue
            if ctype == "text/plain":
                plain += (part.get_payload(decode=True) or b"").decode(
                    part.get_content_charset() or "utf-8", "replace"
                )
            if ctype == "text/html":
                html += (part.get_payload(decode=True) or b"").decode(
                    part.get_content_charset() or "utf-8", "replace"
                )
    else:
        payload = msg.get_payload(decode=True) or b""
        if msg.get_content_type() == "text/plain":
            plain = payload.decode(msg.get_content_charset() or "utf-8", "replace")
        elif msg.get_content_type() == "text/html":
            html = payload.decode(msg.get_content_charset() or "utf-8", "replace")
    if not plain and html: plain = html_to_text(html)
    return plain.strip(), html.strip()

def docx_bytes(subject: str, from_addr: str, date_str: str, body: str) -> bytes:
    """
    Generate a .docx file with email details plus an e-signature section.
    """
    d = Document()
    d.add_heading(subject or "(no subject)", level=1)
    d.add_paragraph(f"From: {from_addr}")
    d.add_paragraph(f"Date: {date_str}")
    d.add_paragraph("")
    d.add_paragraph("Message Body:")
    d.add_paragraph(body or "(no body)")
    d.add_paragraph("")
    d.add_paragraph("\n\n--- E-Signature Section ---\n")
    d.add_paragraph("Signature of Receiver: ________________________________")
    d.add_paragraph("Date: ________________________________")

    bio = BytesIO()
    d.save(bio)
    return bio.getvalue()

def pdf_with_signature(subject: str, from_addr: str, date_str: str, body: str) -> bytes:
    """
    Generate PDF with email content and an interactive signature field.
    """
    # Step 1: create base PDF with reportlab
    tmp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    c = canvas.Canvas(tmp_pdf.name, pagesize=A4)
    width, height = A4

    c.setFont("Helvetica-Bold", 16)
    c.drawString(100, height - 100, subject or "(no subject)")

    c.setFont("Helvetica", 10)
    c.drawString(100, height - 120, f"From: {from_addr}")
    c.drawString(100, height - 135, f"Date: {date_str}")

    c.setFont("Helvetica", 12)
    text_obj = c.beginText(100, height - 160)
    for line in (body or "(no body)").splitlines():
        text_obj.textLine(line.strip())
    c.drawText(text_obj)

    # Placeholder signature box
    c.setFont("Helvetica-Bold", 12)
    c.drawString(100, 150, "Signature of Receiver:")
    c.rect(250, 135, 250, 40)  # visual box for signature

    c.save()

    # Step 2: inject AcroForm signature field
    pdf = PdfReader(tmp_pdf.name)
    page = pdf.pages[0]

    sig_field = IndirectPdfDict(
        FT="/Sig",
        Type="/Annot",
        Subtype="/Widget",
        T="Signature1",
        F=4,
        Rect=[250, 135, 500, 175],
        V=None,
        P=page
    )

    # Ensure Root object exists
    if not getattr(pdf, "Root", None):
        pdf.Root = IndirectPdfDict()

    # Ensure AcroForm exists
    if not getattr(pdf.Root, "AcroForm", None):
        pdf.Root.AcroForm = IndirectPdfDict(Fields=[])

    # Ensure Fields list exists
    if not getattr(pdf.Root.AcroForm, "Fields", None):
        pdf.Root.AcroForm.Fields = []

    # Append field
    pdf.Root.AcroForm.Fields.append(sig_field)
    pdf.Root.AcroForm.SigFlags = 3

    # Ensure page annotations exist
    if not getattr(page, "Annots", None):
        page.Annots = []
    page.Annots.append(sig_field)

    # Step 3: write back into memory
    bio = BytesIO()
    PdfWriter().write(bio, pdf)
    return bio.getvalue()

def as_eml_attachment(raw: bytes) -> Tuple[str, bytes]:
    ts = datetime.utcnow().strftime("%Y%m%d-%H%M%S")
    return (f"original-{ts}.eml", raw)

def safe_json(txt: str) -> Dict[str, Any]:
    try: return json.loads(txt)
    except Exception: pass
    m = re.search(r"\{.*\}", txt, re.DOTALL)
    if m:
        try: return json.loads(re.sub(r",\s*([}\]])", r"\1", m.group(0)))
        except Exception: pass
    return {"category": "general", "subject": None, "reply": txt[:1500]}

def classify_and_draft(subject: str, body: str, sender: str) -> Dict[str, Any]:
    prompt = f"""
Return ONLY JSON with keys: category, subject, reply.

Rules:
- If the email mentions rent, lease, tenancy, landlord, deposit, contract, flat, apartment, house, property, utility bills, repairs, or maintenance → set category="property".
- Otherwise → category="general".

Then draft a short professional reply. If info is missing, ask 2–4 specific follow-up questions.
Subject should typically be 'Re: <original>'.

Original From: {sender}
Original Subject: {subject or '(no subject)'}
Original Body:
\"\"\"{body[:8000]}\"\"\""""
    r = groq.chat.completions.create(
        model=GROQ_MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.2,
        max_tokens=700,
    )
    return safe_json(r.choices[0].message.content.strip())

def build_reply(original: email.message.Message, to_addr: str, subj: str, body: str) -> EmailMessage:
    msg = EmailMessage()
    msg["Subject"] = subj or f"Re: {decode_mime(original.get('Subject')) or '(no subject)'}"
    msg["From"] = f"{SMTP_FROM_NAME} <{SMTP_USER}>"
    msg["To"] = to_addr
    mid = original.get("Message-ID")
    if mid:
        msg["In-Reply-To"] = mid
        refs = original.get("References")
        msg["References"] = f"{refs} {mid}" if refs else mid
    msg.set_content(body)
    return msg

def send_mail(msg: EmailMessage, attachments: Optional[List[Tuple[str, bytes]]] = None):
    if attachments:
        for fn, data in attachments:
            msg.add_attachment(data, maintype="application", subtype="octet-stream", filename=fn)

    if DRY_RUN:
        print(f"[DRY RUN] To={msg['To']} | Subject={msg['Subject']}")
        try:
            body_preview = msg.get_body(preferencelist=('plain'))
            if body_preview:
                print(body_preview.get_content()[:400])
            else:
                print("(no plain text body)")
        except Exception:
            print("(multipart email with attachments — body preview skipped)")
        if attachments:
            print("Attachments:", [a[0] for a in attachments])
        return

    ctx = ssl.create_default_context()
    with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, context=ctx) as s:
        s.login(SMTP_USER, SMTP_PASSWORD)
        s.send_message(msg)

def connect_imap() -> imaplib.IMAP4_SSL:
    ctx = ssl.create_default_context()
    imap = imaplib.IMAP4_SSL(IMAP_HOST, IMAP_PORT, ssl_context=ctx)
    imap.login(IMAP_USER, IMAP_PASSWORD)
    return imap

def mark_seen(imap: imaplib.IMAP4_SSL, mail_id: bytes):
    if not DRY_RUN:
        imap.store(mail_id, "+FLAGS", "\\Seen")

def mail_internal_dt(imap: imaplib.IMAP4_SSL, mail_id: bytes) -> datetime:
    typ, resp = imap.fetch(mail_id, "(INTERNALDATE)")
    if typ != "OK" or not resp or len(resp) < 1:
        return datetime.min.replace(tzinfo=timezone.utc)
    tt = imaplib.Internaldate2tuple(resp[0])
    if tt is None: return datetime.min.replace(tzinfo=timezone.utc)
    return datetime(*tt[:6], tzinfo=timezone.utc)

# ============== Core ==============
def process_one(imap: imaplib.IMAP4_SSL, mail_id: bytes):
    res, data = imap.fetch(mail_id, "(RFC822)")
    if res != "OK": return
    raw = data[0][1]
    msg = email.message_from_bytes(raw)

    from_addr = email.utils.parseaddr(msg.get("From"))[1]
    subj = decode_mime(msg.get("Subject"))
    date_str = decode_mime(msg.get("Date"))

    text, html = extract_text(msg)
    body_in = text or html_to_text(html) or ""

    print(f"\n-- Processing {mail_id.decode()} from {from_addr} | {subj}")

    result = classify_and_draft(subj, body_in, from_addr)
    print("LLM classification result:", result)  # DEBUG PRINT

    category = (result.get("category") or "general").lower().strip()
    reply_subject = result.get("subject") or (f"Re: {subj}" if subj else "Re: Your email")
    reply_text = result.get("reply") or "Thank you for your email."

    # Reply to sender
    reply = build_reply(msg, from_addr, reply_subject, reply_text)
    send_mail(reply)

    # Forward if property
    if category == "property" and FORWARD_TO_EMAIL:
        attachments = [
            ("email_intake.docx", docx_bytes(subj, from_addr, date_str, body_in)),
            ("email_intake.pdf", pdf_with_signature(subj, from_addr, date_str, body_in))
        ]
        if ATTACH_ORIGINAL_EML:
            attachments.append(as_eml_attachment(raw))
        fwd = EmailMessage()
        fwd["Subject"] = f"[Property Intake] {subj or '(no subject)'}"
        fwd["From"] = f"{SMTP_FROM_NAME} <{SMTP_USER}>"
        fwd["To"] = FORWARD_TO_EMAIL
        fwd.set_content(
            "New property-related email attached.\n\n"
            f"From: {from_addr}\nSubject: {subj}\nDate: {date_str}\n"
            "This message was forwarded automatically."
        )
        send_mail(fwd, attachments=attachments)

    mark_seen(imap, mail_id)

def main():
    started_at = datetime.now(timezone.utc)
    imap = connect_imap()
    imap.select(POLL_MAILBOX)

    status, data = imap.search(None, "(UNSEEN)")
    if status != "OK": return
    ids = data[0].split()
    print(f"Unread found: {len(ids)} | only after {started_at.isoformat()}")

    recent_ids = []
    for mid in ids:
        try:
            dt = mail_internal_dt(imap, mid)
            if dt > started_at:
                recent_ids.append(mid)
        except: continue

    print(f"Unread NEW since start: {len(recent_ids)}")

    processed = 0
    for mid in recent_ids:
        try:
            process_one(imap, mid)
            processed += 1
            time.sleep(1.0)
            if processed >= MAX_EMAILS_PER_RUN: break
        except Exception as e:
            print("Error while processing", mid, e)
            traceback.print_exc()
            mark_seen(imap, mid)

    imap.close(); imap.logout()

if __name__ == "__main__":
    print(f"[{datetime.now().isoformat(timespec='seconds')}] Agent start (DRY_RUN={DRY_RUN})")
    main()
    print(f"[{datetime.now().isoformat(timespec='seconds')}] Agent stop")
